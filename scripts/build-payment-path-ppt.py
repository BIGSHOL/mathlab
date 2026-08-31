# -*- coding: utf-8 -*-
"""
토스페이먼츠 카드사 심사용 **결제경로 PPT** 생성기.

## 사양의 출처
토스가 배포하는 공식 가이드 PDF 2종을 그대로 따른다. FAQ 본문이 아니라 **가이드 PDF**가
기준이다 — FAQ 는 "메인-상품-결제-결제창" 4단계로만 적지만, 실제 가이드는 8단계이고
누락 시 보완 요청 사유가 된다.
  · 토스페이먼츠_홈페이지 결제경로 제작 가이드_충전업종용.pdf   (18p, 8단계)
  · 토스페이먼츠_홈페이지 결제경로 제작 가이드_정기결제용.pdf   (16p, 6단계)

파라엑스는 **충전업종(일회성 이용권) + 정기결제(월 구독)** 에 동시에 해당하므로
더 엄격한 충전업종 8단계를 기준으로 삼고, 정기결제 전용 경로를 추가로 싣는다.

## 가이드가 요구하는 순서 (충전업종용 p3)
  ① 가맹점 정보 기재 — 표지에 상호명/사업자번호/URL/**테스트 ID·PW**
  ② 하단정보        — 상호명/사업자번호/대표자명/사업자주소/**상점 전화번호**/**통신판매업신고번호**
  ③④ 약관 및 환불 규정 — **약관도** 캡처(환불만이 아니다)
  ⑤ 로그인 경로     — 비회원 구매 가능하면 생략 가능(우리는 로그인 필수라 **필수**)
  ⑥ 상품선택 / 구매과정
  ⑦ 카드 결제경로   — **2개 카드사 필수**, 비씨/농협은 인증 화면까지
  ⑧ 결제 후 사용처  — 충전한 것이 **어디서 쓰이는지**까지 (충전업종 전용 요건)

가이드 p15 원문: "2개 카드사에 대한 캡처는 필수예요.
                 (결제 완료 화면까지는 캡처 불필요하며 아래처럼 캡처해 주세요.)"
→ 그래서 완료 화면은 **부록**으로 뺀다. 넣어서 감점되진 않지만 필수는 아니다.

스크린샷은 실제 프로덕션 화면을 브라우저로 직접 캡처한 것이다(합성/모형 아님).

실행: python scripts/build-payment-path-ppt.py <스크린샷폴더> <출력.pptx>
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x66, 0x6C, 0x78)
BRAND = RGBColor(0x31, 0x82, 0xF6)
BG = RGBColor(0xF7, 0xF8, 0xFA)
WARN = RGBColor(0xB4, 0x54, 0x09)

# 표지에 기재할 가맹점 정보 — 가이드 p6 의 (1)~(5) 항목 그대로
MERCHANT = [
    ("(1) 상호명", "파라엑스"),
    ("(2) 사업자번호", "496-25-02217"),
    ("(3) URL", "https://mathlab.para-x.co.kr"),
    ("(4) Test ID", "paraxreview"),
    ("(5) Test PW", "Parax2026!"),
]

# (파일명, 단계라벨, 제목, 설명줄들)  — 파일이 없으면 자동 생략
SLIDES = [
    ("02-footer.png", "② 하단정보", "사업자정보 (하단 푸터)",
     ["상호명 · 사업자번호 · 대표자명 · 사업장주소 · 상점 전화번호를 하단에 표기합니다.",
      "통신판매업 신고번호: 신고 진행 중이며, 발급 즉시 동일 영역에 기재할 예정입니다.",
      "이용약관 / 개인정보처리방침 / 취소·환불 정책 링크를 함께 제공합니다."]),
    ("s03-terms.png", "③ 약관", "이용약관",
     ["https://para-x.co.kr/terms.html",
      "서비스 이용 조건, 이용권의 성격(현금화·양도·대여 불가), 해지 조건을 명시합니다.",
      "모든 페이지 하단에서 접근할 수 있습니다."]),
    ("07-refund.png", "④ 환불 규정", "취소·환불 정책",
     ["https://para-x.co.kr/refund.html",
      "무형(디지털) 재화로 배송 개념이 없으며, 서비스 제공기간은 충전일로부터 1년입니다.",
      "결제일로부터 7일 이내 미사용분 전액 환불.",
      "일부 사용 시 사용분을 차감한 잔여분 환불 / 정기결제는 이용 일수 일할 계산."]),
    ("s05-login.png", "⑤ 로그인 경로", "로그인",
     ["https://mathlab.para-x.co.kr/login",
      "학원 단위로 계정을 발급하는 B2B 서비스로, 구매는 로그인 후 진행합니다.",
      "비회원 구매는 제공하지 않으므로 심사용 테스트 계정을 표지에 기재했습니다."]),
    ("01-main.png", "⑥-1 상품선택", "메인 페이지",
     ["https://mathlab.para-x.co.kr",
      "서비스 소개와 기능을 안내합니다. 로그인 후 상단 [결제] 메뉴로 이동합니다."]),
    ("03-products.png", "⑥-2 상품선택", "상품 목록 및 판매가",
     ["https://mathlab.para-x.co.kr/billing",
      "월 결제(정기결제) 3종: Basic ₩80,000 / Pro ₩133,000 / Enterprise ₩280,000 (월)",
      "횟수 결제(일회성 이용권) 3종: 3회 ₩30,000 / 10회 ₩80,000 / 30회 ₩210,000",
      "모든 상품에 판매가를 표기하며, 표기 금액과 실제 결제 금액은 동일합니다."]),
    ("04-checkout.png", "⑥-3 구매과정", "결제 (일회성 이용권)",
     ["https://para-x.co.kr/checkout.html",
      "주문 내용(상품명 · 주문번호 · 결제금액)과 서비스 제공기간을 표시합니다.",
      "결제수단: 신용·체크카드 / 퀵계좌이체 / 토스페이 / 페이코 / 카카오페이 / 네이버페이",
      "결제 서비스 이용약관 및 개인정보 처리 동의 후 [결제하기]를 누릅니다."]),
    ("s07a-card-bc.png", "⑦-1 카드 결제경로", "카드사 인증창 — 비씨카드",
     ["[신용·체크카드] → 카드사 [비씨] 선택 → [결제하기]",
      "비씨카드 ISP/페이북 인증창이 호출됩니다.",
      "이후 카드사 인증을 완료하면 결제가 승인됩니다."]),
    ("s07b-card-nh.png", "⑦-2 카드 결제경로", "카드사 인증창 — 농협카드",
     ["[신용·체크카드] → 카드사 [농협] 선택 → [결제하기]",
      "농협카드 인증창이 호출됩니다(NH Pay 앱카드 / 일반결제 선택).",
      "가이드 요건에 따라 2개 카드사의 인증 화면을 모두 첨부했습니다."]),
    ("05-paywindow.png", "⑦-3 카드 결제경로", "카드사 인증창 — 국민카드 (참고)",
     ["[신용·체크카드] → 카드사 [국민] 선택 → [결제하기]",
      "KB국민카드 인증창 예시입니다."]),
    ("s08-usage.png", "⑧ 결제 후 사용처", "충전한 이용권의 사용 경로",
     ["https://mathlab.para-x.co.kr/exam-analysis",
      "결제로 충전된 이용권은 상단에 잔여 횟수로 표시됩니다.",
      "시험지를 업로드해 분석할 때마다 이용권이 1회씩 차감됩니다.",
      "이용권이 사용되는 유일한 경로이며, 현금화·양도·대여는 불가합니다."]),
    ("08-sub-checkout.png", "정기결제", "정기결제(월 구독) 경로",
     ["https://mathlab.para-x.co.kr/billing → [구독하기]",
      "월 구독은 카드 등록(빌링) 방식으로, 일회성 결제와 화면이 다릅니다.",
      "[카드 등록하고 구독 시작] → 카드사 인증 → 매월 자동 결제 및 이용권 자동 충전",
      "구독은 결제 화면의 [구독 관리]에서 언제든 직접 해지할 수 있습니다."]),
    ("06-success.png", "부록", "결제 완료 및 지급 (참고)",
     ["가이드상 결제 완료 화면은 필수 캡처 대상이 아니나, 상품 지급 근거로 첨부합니다.",
      "결제 승인 즉시 이용권이 학원 계정에 자동 지급됩니다.",
      "결제금액 · 충전 수량 · 주문번호를 표시하고 영수증을 제공합니다."]),
]


def add_text(slide, left, top, width, height, text, size, bold=False,
             color=INK, align=PP_ALIGN.LEFT, spacing=1.0):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.line_spacing = spacing
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = "맑은 고딕"
    return box


def fit(img_path, max_w, max_h):
    with Image.open(img_path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    return int(w * scale), int(h * scale)


def build(shot_dir: Path, out: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ── ① 가맹점 정보 기재 (표지) ── 가이드 p6 형식
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background(); bg.shadow.inherit = False

    add_text(s, Inches(0.9), Inches(0.75), Inches(11.5), Inches(0.6),
             "결제경로 안내", 38, bold=True)
    add_text(s, Inches(0.9), Inches(1.62), Inches(11.5), Inches(0.4),
             "파라엑스 (Para-X) · MathLAB 기출분석", 17, color=MUTED)

    bar = s.shapes.add_shape(1, Inches(0.9), Inches(2.25), Inches(11.5), Inches(0.04))
    bar.fill.solid(); bar.fill.fore_color.rgb = BRAND
    bar.line.fill.background(); bar.shadow.inherit = False

    top = Inches(2.75)
    for label, value in MERCHANT:
        add_text(s, Inches(1.15), top, Inches(2.3), Inches(0.42), label, 16, bold=True, color=BRAND)
        add_text(s, Inches(3.6), top, Inches(8.5), Inches(0.42), value, 16, bold=True)
        top = Emu(int(top) + int(Inches(0.55)))

    add_text(s, Inches(1.15), Inches(6.15), Inches(11.0), Inches(0.9), [
        "· 서비스 유형: 무형(디지털) 재화 — 배송 없음 / 서비스 제공기간: 일회성 이용권 충전일로부터 1년, 구독 1개월 단위",
        "· 결제 유형: 충전업종(일회성 이용권) + 정기결제(월 구독) / 토스페이먼츠 결제위젯 직접 연동(호스팅사 미사용)",
    ], 11.5, color=MUTED, spacing=1.5)

    # ── 본문 ──
    missing = []
    for fname, step, title, notes in SLIDES:
        path = shot_dir / fname
        if not path.exists():
            missing.append(fname)
            continue
        s = prs.slides.add_slide(blank)
        add_text(s, Inches(0.55), Inches(0.32), Inches(4.0), Inches(0.35),
                 step, 13, bold=True, color=BRAND)
        add_text(s, Inches(0.55), Inches(0.62), Inches(12.2), Inches(0.5),
                 title, 25, bold=True)

        img_max_w, img_max_h = int(Inches(8.0)), int(Inches(5.5))
        w, h = fit(path, img_max_w, img_max_h)
        pic = s.shapes.add_picture(
            str(path), Inches(0.55),
            Emu(int(Inches(1.45)) + max(0, (img_max_h - h) // 2)),
            width=Emu(w), height=Emu(h))
        pic.line.color.rgb = RGBColor(0xDD, 0xE1, 0xE6)
        pic.line.width = Pt(0.75)

        add_text(s, Inches(8.85), Inches(1.5), Inches(3.95), Inches(5.0),
                 [f"• {n}" for n in notes], 12.5, spacing=1.6)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    n = len(prs.slides._sldIdLst)
    print(f"생성 완료: {out}  (슬라이드 {n}장)")
    if missing:
        print("!! 캡처 누락으로 빠진 슬라이드: " + ", ".join(missing))


if __name__ == "__main__":
    build(Path(sys.argv[1]), Path(sys.argv[2]))
